from __future__ import annotations
from io import BytesIO
from typing import Any
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def workbook_bytes(rows: list[dict[str, Any]], domain: str) -> bytes:
    wb=Workbook(); ws=wb.active; ws.title="Workforce"
    headers=["Domain","Name","Title","Member","Source","Confidence","Observation"]
    ws.append(headers)
    for idx,row in enumerate(rows[:5000],1):
        ws.append([domain,row.get("name"),row.get("title"),row.get("member"),row.get("source"),row.get("confidence"),row.get("record",idx)])
    for i,width in enumerate([28,28,48,34,24,14,15],1):
        ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes="A2"
    meta=wb.create_sheet("README")
    meta.append(["Field","Value"])
    meta.append(["Domain",domain])
    meta.append(["Purpose","Public/provider-returned workforce observations"])
    meta.append(["Collection policy","No LinkedIn page scraping; official APIs/provider-returned rows only"])
    meta.column_dimensions["A"].width=24; meta.column_dimensions["B"].width=100
    out=BytesIO(); wb.save(out); return out.getvalue()
