from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape


def _zip_xml(files: dict[str, str | bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def docx(title, scan):
    paras = [
        title,
        f"Risk: {scan.get('risk_score', 0)}/100 · Verdict: {scan.get('verdict', 'unknown')}",
        str(scan.get('executive_summary', '')),
    ]
    for f in scan.get('findings', [])[:100]:
        paras.append(f"{f.get('id')}: {f.get('name')} | {f.get('severity')} | confidence {f.get('confidence')} | {f.get('what_found')} | {f.get('interpretation')}")
    body = ''.join(f'<w:p><w:r><w:t xml:space="preserve">{escape(str(p))}</w:t></w:r></w:p>' for p in paras)
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f'<w:body>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr>'
                '</w:body></w:document>')
    content = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
               '<Default Extension="xml" ContentType="application/xml"/>'
               '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
               '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>')
    return _zip_xml({'[Content_Types].xml': content, '_rels/.rels': rels, 'word/document.xml': document})


def _rows(scan: dict) -> dict[str, list[list[str]]]:
    health = scan.get('health') or scan.get('domain_health') or {}
    checks = []
    if isinstance(health, dict):
        for c in health.get('checks', [])[:1000]:
            checks.append([
                c.get('name', ''), c.get('status', c.get('state', '')), c.get('detail', ''),
                c.get('source', ''), c.get('confidence', ''), c.get('module', ''),
                ', '.join(str(x) for x in (c.get('pivots') or [])),
            ])
    findings = []
    for f in scan.get('findings', [])[:1000]:
        findings.append([f.get('id', ''), f.get('name', ''), f.get('severity', ''), f.get('confidence', ''), f.get('what_found', ''), f.get('interpretation', '')])
    assets = []
    for kind, values in (scan.get('iocs') or {}).items():
        if isinstance(values, list):
            assets.extend([[kind, str(v)] for v in values[:2000]])
    evidence = []
    for e in (scan.get('evidence') or scan.get('raw_evidence') or [])[:2000]:
        if isinstance(e, dict):
            evidence.append([
                e.get('id', e.get('evidence_id', '')), e.get('type', ''), e.get('source', ''),
                e.get('timestamp', e.get('observed_at', '')), e.get('asset_id', ''), e.get('sha256', ''), e.get('scan_id', ''),
            ])
        else:
            evidence.append(['', type(e).__name__, '', '', '', '', ''])
    timeline = []
    for e in (scan.get('timeline') or [])[:2000]:
        if isinstance(e, dict):
            timeline.append([
                e.get('timestamp', e.get('observed_at', e.get('date', ''))),
                e.get('type', e.get('event_type', '')), e.get('asset_id', e.get('entity_id', '')),
                e.get('description', e.get('summary', '')), e.get('source', ''),
            ])
        else:
            timeline.append(['', '', '', str(e), ''])
    osint = []
    adv = scan.get('advanced_osint') or {}
    if isinstance(adv, dict):
        for c in adv.get('categories', [])[:500]:
            if isinstance(c, dict):
                osint.append([c.get('id', c.get('name', '')), c.get('name', ''), c.get('observed_fields', ''), c.get('direct_link_count', ''), c.get('execution', '')])
    ai = scan.get('silent_ai') or {}
    summary = [
        ['Report', 'DOMINUS ULTIMATE Analyst Report'],
        ['Title', scan.get('title', '')],
        ['Target', scan.get('domain', scan.get('target', ''))],
        ['Risk', scan.get('risk_score', 0)],
        ['Verdict', scan.get('verdict', 'unknown')],
        ['Confidence', scan.get('confidence', scan.get('confidence_score', ''))],
        ['Executive Summary', scan.get('executive_summary', '')],
        ['Silent AI', ai.get('status', 'embedded')],
        ['Silent AI Coverage', ai.get('coverage_percent', '')],
    ]
    return {
        'Summary': summary,
        'Health Recon': [['Check', 'Status', 'Detail', 'Source', 'Confidence', 'Module', 'Pivots']] + checks,
        'Findings': [['ID', 'Name', 'Severity', 'Confidence', 'What Found', 'Interpretation']] + findings,
        'Assets': [['Type', 'Value']] + assets,
        'Evidence': [['Evidence ID', 'Type', 'Source', 'Timestamp', 'Asset ID', 'SHA-256', 'Scan ID']] + evidence,
        'Timeline': [['Timestamp', 'Type', 'Asset', 'Description', 'Source']] + timeline,
        'OSINT': [['Category ID', 'Category', 'Observed Fields', 'Direct Links', 'Execution']] + osint,
    }


def _col_letter(n: int) -> str:
    out = ''
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _sheet_xml(rows: list[list[str]]) -> str:
    xml_rows = []
    for r_idx, row in enumerate(rows or [['']], 1):
        cells = []
        for c_idx, value in enumerate(row, 1):
            ref = f'{_col_letter(c_idx)}{r_idx}'
            val = escape(str('' if value is None else value))
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{val}</t></is></c>')
        xml_rows.append(f'<row r="{r_idx}">' + ''.join(cells) + '</row>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetViews><sheetView workbookViewId="0"/></sheetViews><sheetData>'
            + ''.join(xml_rows) + '</sheetData></worksheet>')


def xlsx(title, scan):
    sheets = _rows(scan)
    names = list(sheets)
    sheet_defs, rel_defs, overrides = [], [], []
    files: dict[str, str | bytes] = {}
    for idx, name in enumerate(names, 1):
        safe = escape(name, {'"': '&quot;'})
        sheet_defs.append(f'<sheet name="{safe}" sheetId="{idx}" r:id="rId{idx}"/>')
        rel_defs.append(f'<Relationship Id="rId{idx}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{idx}.xml"/>')
        overrides.append(f'<Override PartName="/xl/worksheets/sheet{idx}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
        files[f'xl/worksheets/sheet{idx}.xml'] = _sheet_xml(sheets[name])
    workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<workbookPr/><sheets>' + ''.join(sheet_defs) + '</sheets></workbook>')
    root_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wb_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + ''.join(rel_defs) + '</Relationships>')
    content = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
               '<Default Extension="xml" ContentType="application/xml"/>'
               '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
               '</Types>')
    # Insert all worksheet overrides before closing Types.
    content = content.replace('</Types>', ''.join(overrides) + '</Types>')
    files.update({
        '[Content_Types].xml': content,
        '_rels/.rels': root_rels,
        'xl/workbook.xml': workbook,
        'xl/_rels/workbook.xml.rels': wb_rels,
    })
    return _zip_xml(files)


def pptx(title, scan):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        pres = Presentation(); pres.slide_width = Inches(13.333); pres.slide_height = Inches(7.5)
        slides = [
            ('DOMINUS ULTIMATE Analyst Report', [title, f"Risk {scan.get('risk_score', 0)}/100 · {scan.get('verdict', 'unknown')}"]),
            ('Executive Assessment', [str(scan.get('executive_summary', ''))]),
            ('Key Findings', [f"{f.get('name')} — {f.get('severity')} — confidence {f.get('confidence')}" for f in scan.get('findings', [])[:15]]),
        ]
        for heading, lines in slides:
            slide = pres.slides.add_slide(pres.slide_layouts[5])
            if slide.shapes.title: slide.shapes.title.text = heading
            frame = slide.shapes.add_textbox(Inches(.8), Inches(1.5), Inches(11.8), Inches(5)).text_frame
            for i, line in enumerate(lines):
                par = frame.paragraphs[0] if i == 0 else frame.add_paragraph(); par.text = str(line); par.font.size = Pt(18)
        out = io.BytesIO(); pres.save(out); return out.getvalue()
    except Exception:
        return docx(title, scan)
